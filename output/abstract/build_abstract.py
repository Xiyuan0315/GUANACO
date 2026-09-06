from pathlib import Path

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn


OUT = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = sec.bottom_margin = Mm(20)
sec.left_margin = sec.right_margin = Mm(22)

for name in ('Normal', 'Title', 'Heading 1'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        style.element.rPr.rFonts.attrib.pop(qn('w:' + attr), None)
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(6)

doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.keep_with_next = True

def labeled(label, body, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(label + '\n').bold = True
    p.add_run(body)
    p.paragraph_format.widow_control = True
    return p

labeled('Title', 'GUANACO for interactive exploration of single cell and spatial multiomics data', 'Title')
labeled('Authors', 'Zhang X, Kuddus M, Xia Q, Hu Y, Chen P')
labeled('KI department', '[Insert KI department], Karolinska Institutet, Stockholm, Sweden')
labeled('Affiliation with Region Stockholm (if applicable)', '[Insert affiliation, or state Not applicable]')

sections = {
    'Background': (
        'Single-cell and spatial technologies reveal how cells differ in their molecular '
        'profiles and tissue locations. Interpreting these data often requires switching '
        'between specialised software and writing code, limiting direct exploration by '
        'experimental and clinical researchers. We developed GUANACO (Graphical Unified '
        'Analysis and Navigation of Cellular Omics) to make these datasets accessible '
        'through coordinated interactive visualisation.'
    ),
    'Methods': (
        'GUANACO is an open-source Python platform with a browser dashboard and interfaces '
        'for computational notebooks. It reads standard single-cell and multiomics data '
        'formats and connects plots through shared cell, sample or feature identifiers. '
        'Selecting a population or gene updates linked views, while selected identifiers '
        'can be retrieved for further analysis. Molecular measurements from the same cells '
        'can be compared directly; measurements from different cells can be summarised '
        'across shared biological groups. Data are loaded as needed for the active view.'
    ),
    'Results': (
        'The platform brings together cell population maps, gene expression summaries, '
        'tissue images and genome tracks in a configurable interface. It also displays '
        'precomputed results describing cell-state trajectories, spatial neighbourhoods '
        'and predicted cell-to-cell signalling. Included examples cover blood immune '
        'cells, spatial tissue data and a public chronic lymphocytic leukaemia cohort '
        'of 200 patients. The latter combines gene expression, DNA methylation, genomic '
        'alterations and experimental drug responses, preserving the actual patient '
        'coverage of each measurement type. These workflows support exploration of '
        'molecular patterns alongside cell annotations and clinical characteristics.'
    ),
    'Conclusions': (
        'GUANACO provides a common visual workspace for exploring cellular and clinical '
        'multiomics data. By connecting interactive figures with programmable analysis, '
        'it is designed to support hypothesis generation and collaboration between '
        'biomedical researchers and computational scientists.'
    ),
}
for label, body in sections.items():
    labeled(label, body)

for p in doc.paragraphs:
    for border in list(p._p.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            run._element.rPr.rFonts.attrib.pop(qn('w:' + attr), None)

doc.core_properties.title = 'GUANACO abstract for Collaboration in Science 2026'
doc.core_properties.subject = 'Conference abstract'
doc.core_properties.author = ''
doc.core_properties.keywords = 'GUANACO, multiomics, visualisation'
path = OUT / 'GUANACO_Abstract_Collaboration_in_Science_2026.docx'
doc.save(path)
print(path)
print('Body word count:', sum(len(s.split()) for s in sections.values()))
